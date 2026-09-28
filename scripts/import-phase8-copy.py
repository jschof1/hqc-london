"""Extract customer-facing sections from Pedro's approved Batch 1–5 DOCX files.

Usage: python3 scripts/import-phase8-copy.py /path/to/batch1.docx ... /path/to/batch5.docx
The release pack's Section 15 CTA amendments are applied separately in the site.
"""

import json
import re
import sys
from pathlib import Path

from docx import Document


PAGE_SLUGS = {
    (1, 1): 'home',
    (1, 2): 'property-facilities-cleaning-london',
    (1, 3): 'office-cleaning-london',
    (1, 4): 'home-cleaning-london',
    (2, 1): 'essential-clean-london',
    (2, 2): 'deep-cleaning-london',
    (2, 3): 'end-of-tenancy-cleaning-london',
    (2, 4): 'move-in-cleaning-london',
    (2, 5): 'post-construction-cleaning-london',
    (3, 1): 'communal-area-cleaning-london',
    (3, 2): 'estate-letting-agent-cleaning-london',
    (3, 3): 'move-in-void-property-cleaning-london',
    (3, 4): 'portfolio-cleaning-support-london',
    (3, 5): 'commercial-cleaning-london',
    (3, 6): 'commercial-deep-cleaning-london',
    (4, 1): 'complex-property-cleaning-london',
    (4, 2): 'how-we-work',
    (4, 3): 'cleaning-prices-london',
    (4, 4): 'why-choose-hqc',
    (4, 5): 'about-high-quality-clean',
    (5, 2): 'reviews-case-studies',
    (5, 5): 'areas-we-cover',
    (5, 6): 'cleaning-services-east-london',
}


def extract(path: Path, batch: int):
    pages = {}
    current = None
    section = None
    in_hero = False
    for paragraph in Document(path).paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        style = paragraph.style.name
        match = re.match(r'^PAGE (\d+) - ', value) if style == 'Heading 1' else None
        if match:
            slug = PAGE_SLUGS.get((batch, int(match.group(1))))
            current = {'title': '', 'lead': [], 'heroItems': [], 'sections': [], 'source': f'Batch {batch}, page {match.group(1)}'} if slug else None
            if slug:
                pages[slug] = current
            section = None
            in_hero = True
            continue
        if style == 'Heading 1':
            current = None
            continue
        if current is None:
            continue
        if style == 'Heading 2':
            in_hero = value == 'Hero'
            section = None if in_hero else {'heading': value, 'blocks': []}
            if section:
                current['sections'].append(section)
            continue
        if in_hero:
            if style == 'Heading 3' and value.startswith('H1 - '):
                current['title'] = value[5:]
            elif style == 'Normal' and not re.match(r'^(Primary CTA|Secondary CTA|Trust line|Supporting trust line):', value, re.I):
                current['lead'].append(value)
            elif style.startswith('List'):
                current['heroItems'].append(value)
            continue
        if not section or re.match(r'^(Primary CTA|Secondary CTA|CTA|Trust line|Supporting line):', value, re.I):
            continue
        if value.startswith('[') and value.endswith(']'):
            continue
        kind = 'h3' if style == 'Heading 3' else 'li' if style.startswith('List') else 'p'
        section['blocks'].append({'kind': kind, 'text': value})
    for page in pages.values():
        for section in page['sections']:
            grouped = []
            for block in section['blocks']:
                if block['kind'] == 'li':
                    if not grouped or grouped[-1]['kind'] != 'ul':
                        grouped.append({'kind': 'ul', 'items': []})
                    grouped[-1]['items'].append(block['text'])
                else:
                    grouped.append(block)
            section['blocks'] = grouped
    return pages


def extract_faq(path: Path):
    groups = []
    active = False
    group = None
    question = None
    for paragraph in Document(path).paragraphs:
        value = paragraph.text.strip()
        if paragraph.style.name == 'Heading 1' and value.startswith('PAGE 1 - CLEANING FAQs'):
            active = True
            continue
        if paragraph.style.name == 'Heading 1' and active:
            break
        if not active or not value:
            continue
        if paragraph.style.name == 'Heading 2':
            group = {'title': value, 'items': []} if value not in ('Hero', 'Still Not Sure?') else None
            if group:
                groups.append(group)
            question = None
        elif paragraph.style.name == 'Heading 3' and value.endswith('?') and group:
            question = {'question': value, 'answer': ''}
            group['items'].append(question)
        elif question and paragraph.style.name == 'Normal':
            question['answer'] += (' ' if question['answer'] else '') + value
    return [group for group in groups if group['items']]


if __name__ == '__main__':
    if len(sys.argv) != 6:
        raise SystemExit(__doc__)
    result = {}
    for batch, filename in enumerate(sys.argv[1:], 1):
        result.update(extract(Path(filename), batch))
    missing = set(PAGE_SLUGS.values()) - result.keys()
    if missing:
        raise SystemExit(f'Missing controlled pages: {sorted(missing)}')
    output = Path(__file__).resolve().parent.parent / 'src/data/phase8-copy.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Wrote {len(result)} approved page bodies to {output}')
    faq = extract_faq(Path(sys.argv[5]))
    faq_output = output.with_name('phase8-faq.json')
    faq_output.write_text(json.dumps(faq, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Wrote {sum(len(group["items"]) for group in faq)} approved FAQs to {faq_output}')
